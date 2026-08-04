#!/usr/bin/env python3
"""Create the narrative GOEI reading document from its controlled Markdown source."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_integral_dossier import BLUE, DARK, MUTED, ROOT, configure_document


SOURCE = ROOT / "12_Tesis" / "GOEI_Relato_Integral_Proyecto_2026-08-04.md"
OUTPUT = ROOT / "12_Tesis" / "GOEI_Documento_de_Lectura_Proyecto_2026-08-04.docx"
FIGURES = ROOT / "12_Tesis" / "figuras" / "metodologia"

VISUALS = {
    "Nota de lectura": [
        ("01_mapa_mental_metodologia_goei.png", "Mapa mental de la metodología sustantiva del GOEI")
    ],
    "8. B1 y D-008: el núcleo universal de procesos": [
        ("02_operaciones_goei_p1_p9.png", "Arquitectura operacional universal y lenguaje de modelado")
    ],
    "11. La ampliación GTRM": [
        ("03_observacion_interpretacion_goei.png", "Cadena metodológica desde la observación hasta la interpretación")
    ],
    "12. B4: el punto exacto donde está detenido el trabajo científico": [
        ("04_construccion_metodologica_b1_b10.png", "Construcción metodológica del estándar GOEI")
    ],
}


def clean(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text


def add_cover(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(80)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("GOEI · DOCUMENTO DE LECTURA")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Historia, construcción y estado actual del proyecto")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = DARK
    title.paragraph_format.space_after = Pt(12)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Un relato continuo desde el origen del GOEI hasta sus decisiones, pendientes y próximos entregables")
    run.font.size = Pt(14)
    run.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(30)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Corte documental: 4 de agosto de 2026\nDocumento informativo · No normativo")
    run.font.size = Pt(10.5)
    run.font.color.rgb = MUTED
    meta.paragraph_format.line_spacing = 1.35
    doc.add_page_break()


def new_numbering_id(doc: Document) -> int:
    style = doc.styles["List Number"]
    base_num_id = int(style._element.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val")))
    numbering = doc.part.numbering_part.element
    base_num = numbering.xpath(f'./w:num[@w:numId="{base_num_id}"]')[0]
    abstract_id = int(base_num.xpath("./w:abstractNumId")[0].get(qn("w:val")))
    concrete = numbering.add_num(abstract_id)
    concrete.add_lvlOverride(ilvl=0).add_startOverride(1)
    return int(concrete.numId)


def add_numbered_paragraph(doc: Document, text: str, num_id: int, compact: bool = False) -> None:
    paragraph = doc.add_paragraph(clean(text), style="List Number")
    if compact:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def add_visuals(doc: Document, heading_text: str) -> None:
    for filename, caption in VISUALS.get(heading_text, []):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run().add_picture(str(FIGURES / filename), width=Inches(6.75))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        run = cap.add_run(caption)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def build() -> None:
    doc = Document()
    configure_document(doc)
    doc.sections[0].header.paragraphs[0].text = "GOEI · Historia, construcción y estado actual · Documento de lectura"
    add_cover(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    skip_first_title = True
    current_num_id = None
    compact_numbering = False
    for original in lines:
        line = original.strip()
        if not line:
            current_num_id = None
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            current_num_id = None
            if skip_first_title:
                skip_first_title = False
                continue
            level = min(3, len(heading.group(1)))
            heading_text = clean(heading.group(2))
            compact_numbering = heading_text.startswith("Anexo ")
            doc.add_heading(heading_text, level=level)
            add_visuals(doc, heading_text)
            continue
        if re.match(r"^\d+[.)]\s+", line):
            content = re.sub(r"^\d+[.)]\s+", "", line)
            if current_num_id is None:
                current_num_id = new_numbering_id(doc)
            add_numbered_paragraph(doc, content, current_num_id, compact=compact_numbering)
            continue
        current_num_id = None
        if re.match(r"^[-*+]\s+", line):
            content = re.sub(r"^[-*+]\s+", "", line)
            doc.add_paragraph(clean(content), style="List Bullet")
            continue
        p = doc.add_paragraph(clean(line))
        p.paragraph_format.widow_control = True

    props = doc.core_properties
    props.title = "GOEI — Historia, construcción y estado actual"
    props.subject = "Documento narrativo de lectura del proyecto GOEI"
    props.author = "GOEI Technical Office — Codex"
    props.keywords = "GOEI, historia, decisiones, pendientes, tesis, estándar"
    props.comments = "Síntesis narrativa basada en los registros controlados del repositorio."
    props.created = datetime(2026, 8, 4)
    doc.save(OUTPUT)
    print(f"created={OUTPUT}")


if __name__ == "__main__":
    build()
