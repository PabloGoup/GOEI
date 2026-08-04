#!/usr/bin/env python3
"""Build a linear, provenance-preserving GOEI project dossier."""

from __future__ import annotations

import hashlib
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "12_Tesis" / "GOEI_Dossier_Integral_Proyecto_2026-08-04.docx"

INCLUDE_ROOTS = [
    "00_Gobernanza",
    "01_Estado_del_arte",
    "02_Bibliografia",
    "03_Taxonomia_negocios",
    "04_Procesos_gastronomicos",
    "05_Catalogo_KPI",
    "06_Normalizacion",
    "07_Ponderacion_y_formula",
    "08_Interpretacion",
    "09_Manual_tecnico_general",
    "10_Manual_piloto",
    "11_Validacion",
    "12_Tesis",
    "13_Publicaciones",
    "15_Mantenimiento",
]

TEXT_EXTENSIONS = {".md", ".txt", ".json", ".yaml", ".yml", ".csv", ".bib"}
EXCLUDED_PARTS = {"cuarentena", "__pycache__"}
EXCLUDED_NAMES = {"knowledge_graph.registry.json", OUTPUT.name}

BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(31, 58, 95)
MUTED = RGBColor(89, 89, 89)
LIGHT = "E8EEF5"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after, color in [
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 11.5, 9, 4, DARK),
        ("Heading 4", 10.5, 7, 3, DARK),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.38)
        style.paragraph_format.first_line_indent = Inches(-0.19)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.text = "GOEI · Dossier integral del proyecto · Documento de lectura"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED


def collect_sources() -> list[Path]:
    files: list[Path] = []
    for root_name in INCLUDE_ROOTS:
        base = ROOT / root_name
        if not base.exists():
            continue
        for path in base.rglob("*"):
            if not path.is_file() or path.suffix.lower() not in TEXT_EXTENSIONS:
                continue
            if path.name in EXCLUDED_NAMES or EXCLUDED_PARTS.intersection(path.parts):
                continue
            files.append(path)
    for name in ("README.md", "CHANGELOG.md", "AGENTS.md"):
        path = ROOT / name
        if path.exists():
            files.append(path)
    return sorted(set(files), key=source_sort_key)


def source_sort_key(path: Path) -> tuple:
    rel = path.relative_to(ROOT).as_posix()
    priority = {
        "README.md": -30,
        "AGENTS.md": -29,
        "CHANGELOG.md": -28,
        "00_Gobernanza/Constitucion_GOEI.md": -27,
        "00_Gobernanza/ESTADO_MAESTRO.md": -26,
    }
    if rel in priority:
        return (priority[rel], rel)
    return (0, rel.casefold())


def add_cover(doc: Document, count: int) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("GOEI · GASTRONOMIC OPERATIONAL EFFICIENCY INDEX")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Dossier integral del proyecto")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(28)
    run.font.color.rgb = DARK
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Historia, decisiones, arquitectura, avances, evidencia documental y pendientes")
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(28)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Corte documental: 4 de agosto de 2026\n{count} archivos fuente consolidados\nEstado: dossier informativo, no normativo").font.size = Pt(11)
    meta.paragraph_format.line_spacing = 1.35

    doc.add_page_break()


def add_status_summary(doc: Document) -> None:
    doc.add_heading("1. Cómo leer este dossier", level=1)
    doc.add_paragraph(
        "Este documento reúne linealmente los archivos relevantes del repositorio GOEI. La síntesis inicial permite comprender el estado del proyecto; después se incorpora la transcripción de las fuentes con su ruta, tamaño y huella SHA-256. La transcripción no reemplaza los originales ni modifica su autoridad, versión o estado."
    )

    p = doc.add_paragraph()
    r = p.add_run("Regla de autoridad. ")
    r.bold = True
    p.add_run("Las decisiones científicas pertenecen al Comité Científico; las resoluciones directivas aprueban o devuelven; la Oficina Técnica integra, versiona y audita. Un documento candidato, auditoría o resumen no equivale a una decisión científica aprobada.")

    doc.add_heading("2. Resumen ejecutivo del estado actual", level=1)
    summary = [
        ("Gobernanza y repositorio", "Arquitectura documental, Constitución, reglas Git, convenciones, registros maestros, control de versiones, releases, baselines y CHANGELOG se encuentran implantados."),
        ("Ontología", "D-006 permanece congelada como Ontology Baseline v0.1. Toda ampliación requiere Issue formal y una nueva decisión/versionado; el candidato v0.2.0 no sustituye la baseline."),
        ("Procesos B1", "D-008 fue ratificada y la baseline GOEI-BSL-PROC-0001 v0.1.0 quedó congelada con P1–P9. La investigación gastronómica concreta no se incorpora por inferencia."),
        ("Taxonomía B2", "D-017 fue ratificada y GOEI-BSL-TAX-0001 v0.1.0 quedó congelada: modelo facetado, unidad Establecimiento, categoría como combinación y facetas FX-1, FX-2 y FX-3, sin valores concretos todavía."),
        ("Universo KPI B3", "D-018 fue ratificada y GOEI-BSL-KPI-C-0001 v0.1.0 congeló 31 candidatos evaluables. C-032 Balanced Scorecard se conserva como marco contextual y no como KPI."),
        ("GTRM", "D-009 a D-016 separan observación, medición, valor KPI, referencias, benchmark, función de puntuación, normalización, puntuación e interpretación. La GTRM actúa sobre valores KPI y su aplicación produce la puntuación normalizada; no existe doble normalización."),
        ("B4", "La preparación técnica está completa mediante GOEI-AUDIT-0007. Falta el dictamen científico D-019 sobre método, criterios y admisibilidad de los 31 candidatos."),
        ("Tesis y publicación", "Existen arquitectura, matriz de requisitos y esqueletos documentales, pero no una tesis científicamente terminada ni un manuscrito publicable. Los 64 requisitos directorales aún no tienen estado Aprobado o Publicado."),
        ("Automatización", "La dirección automática se encuentra pausada por instrucción del Director General. El reinicio deberá abrir una conversación nueva con Claude y utilizar paquetes de contexto compactos."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.65)
    table.columns[1].width = Inches(5.1)
    for idx, label in enumerate(("Área", "Estado")):
        cell = table.rows[0].cells[idx]
        cell.text = label
        set_cell_shading(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for area, state in summary:
        cells = table.add_row().cells
        cells[0].text = area
        cells[1].text = state
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for cell in cells:
            set_cell_margins(cell)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    doc.add_heading("3. Decisiones y conclusiones vigentes", level=1)
    decisions = [
        "GOEI-DIR-0004 delega la dirección operativa y permite avanzar sin validaciones intermedias del usuario dentro de la autoridad ya conferida.",
        "D-006 es la ontología conceptual congelada; cualquier modificación exige Issue formal.",
        "D-008 establece la arquitectura universal de procesos y habilitó B2.",
        "D-017 establece la estructura taxonómica universal y habilitó B3.",
        "D-018 acepta exclusivamente el universo candidato B3; no aprueba ningún KPI.",
        "D-009–D-016 constituyen el marco científico GTRM ratificado, con reservas e Issues todavía activos para implementaciones concretas.",
        "GOEI-DIR-RES-0001, 0002 y 0003 congelaron respectivamente las baselines de procesos, taxonomía y candidatos KPI.",
    ]
    for item in decisions:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("4. Pendientes y ruta crítica", level=1)
    pending = [
        "Reanudar B4 con conversación nueva de Claude y paquete compacto; resolver método, criterios, 31 evaluaciones, discrepancias y D-019.",
        "Ejecutar B5: especificaciones de medición, variables, unidades, periodos y condiciones de referencia para KPI admisibles.",
        "Ejecutar B6: funciones GTRM/puntuación y normalización, sin doble transformación.",
        "Ejecutar B7: ponderación, agregación y reglas de no compensación mediante decisión científica.",
        "Ejecutar B8: interpretación, estados de validez e incertidumbre.",
        "Ejecutar B9: formularios, tablas manuales y equivalencia papel–software.",
        "Ejecutar B10: Manual Técnico y conocimiento normativo integrado.",
        "Resolver fundamentos provisionales D-003/D-005 e Issues 006, 007, 016 y 017 de la fase A.",
        "Construir candidato completo del estándar, diseñar y ejecutar validación por expertos/casos reales, y resolver ISS-041.",
        "Completar tesis y publicaciones solo después de disponer de resultados científicos y validación; resolver ISS-039 e ISS-040.",
        "Mantener abiertos y selectivos ISS-055–058, 008, 015, 043, 053 y 054 hasta cumplir sus condiciones documentadas.",
    ]
    for item in pending:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("5. Entregables finales previstos", level=1)
    deliverables = [
        "Estándar GOEI validado y versionado.",
        "Constitución, Body of Knowledge, ontología, taxonomía y modelos de procesos consolidados.",
        "Catálogo de KPI admisibles, especificaciones de medición, GTRM, normalización, fórmula, reglas críticas e interpretación aprobadas.",
        "Formularios manuales, tablas de consulta, Manual Técnico General y manuales piloto.",
        "Knowledge Graph y componentes preparados para software sin convertir el software en fuente normativa.",
        "Expediente de validación científica y operacional.",
        "Tesis terminada derivada del estándar y uno o más manuscritos preparados para publicación.",
    ]
    for item in deliverables:
        doc.add_paragraph(item, style="List Bullet")


def add_source_index(doc: Document, files: list[Path]) -> None:
    doc.add_page_break()
    doc.add_heading("6. Índice de fuentes incorporadas", level=1)
    doc.add_paragraph(f"Total: {len(files)} archivos. El orden sigue la arquitectura del repositorio y, dentro de cada carpeta, la ruta alfabética.")
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(0.55)
    table.columns[1].width = Inches(6.2)
    for idx, label in enumerate(("N.º", "Ruta fuente")):
        cell = table.rows[0].cells[idx]
        cell.text = label
        set_cell_shading(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, path in enumerate(files, 1):
        rel = path.relative_to(ROOT).as_posix()
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = rel
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for cell in cells:
            set_cell_margins(cell, top=40, bottom=40)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)


def clean_inline_markdown(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text


def add_source_content(doc: Document, path: Path, number: int) -> None:
    rel = path.relative_to(ROOT).as_posix()
    raw = path.read_text(encoding="utf-8", errors="replace")
    digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()

    doc.add_page_break()
    heading = doc.add_heading(f"Fuente {number}: {rel}", level=1)
    heading.paragraph_format.keep_with_next = True
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(10)
    r = meta.add_run(f"Ruta: {rel}\nTamaño: {len(raw.encode('utf-8'))} bytes\nSHA-256: {digest}")
    r.font.name = "Consolas"
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    in_code = False
    for original in raw.splitlines():
        line = original.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if not line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            continue
        if in_code or path.suffix.lower() in {".json", ".yaml", ".yml", ".csv", ".bib"}:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)
            continue
        match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if match:
            source_level = len(match.group(1))
            level = min(4, source_level + 1)
            doc.add_heading(clean_inline_markdown(match.group(2)), level=level)
            continue
        if re.match(r"^\s*[-*+]\s+", line):
            content = re.sub(r"^\s*[-*+]\s+", "", line)
            doc.add_paragraph(clean_inline_markdown(content), style="List Bullet")
            continue
        if re.match(r"^\s*\d+[.)]\s+", line):
            content = re.sub(r"^\s*\d+[.)]\s+", "", line)
            doc.add_paragraph(clean_inline_markdown(content), style="List Number")
            continue
        if line.lstrip().startswith("|"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(7.2)
            continue
        if line.lstrip().startswith(">"):
            p = doc.add_paragraph(clean_inline_markdown(line.lstrip()[1:].strip()))
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_after = Pt(5)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = MUTED
            continue
        doc.add_paragraph(clean_inline_markdown(line))


def main() -> None:
    files = collect_sources()
    doc = Document()
    configure_document(doc)
    add_cover(doc, len(files))
    add_status_summary(doc)
    add_source_index(doc, files)
    doc.add_page_break()
    doc.add_heading("7. Transcripción lineal de fuentes", level=1)
    doc.add_paragraph("Desde este punto se reproducen los archivos seleccionados. Cada fuente comienza en una página nueva y conserva su procedencia mediante ruta y SHA-256.")
    for index, path in enumerate(files, 1):
        add_source_content(doc, path, index)

    core = doc.core_properties
    core.title = "GOEI — Dossier integral del proyecto"
    core.subject = "Historia, decisiones, arquitectura, avances y pendientes"
    core.author = "GOEI Technical Office — Codex"
    core.keywords = "GOEI, estándar, gobernanza, decisiones, trazabilidad, tesis"
    core.comments = "Dossier informativo generado desde fuentes versionadas; no modifica su autoridad."
    core.created = datetime(2026, 8, 4)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"sources={len(files)}")


if __name__ == "__main__":
    main()
